"""
Create a new PowerPoint file with a single slide animating a door image swinging inwards.

Usage:
    door-animate.py <door-image-path> --frames=<n> --frame-time=<ms> [--initial-delay=<ms>]
    door-animate.py (-h | --help)
    door-animate.py --version

Options:
    -h --help                 Show this help message.
    --frames=<n>              Number of frames.
    --frame-time=<ms>         Frame time in ms.
    --initial-delay=<ms>      Initial animation delay in ms [default: 0].
"""

import io
import tempfile

from docopt import docopt
import cv2
import numpy as np
from lxml import etree
from pptx import Presentation
from pptx.util import Cm
from pptx.oxml.ns import nsuri

# the "p" xml namespace
p = "{" + nsuri('p') + "}"

template_dir = "door-animate"


def find_target(element, attrib_name):
    """"""
    target = element.find(f".//{p}*[@{attrib_name}='']")
    del target.attrib[attrib_name]
    return target


def main(door_image_path, frames, frame_time, initial_delay=0) -> int:
    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)
    left = top = Cm(1)

    # Load image
    img = cv2.imread(door_image_path, cv2.IMREAD_UNCHANGED)
    h, w = img.shape[:2]

    # Source points (corners of the original image)
    pts_src = np.float32([[0, 0], [w, 0], [w, h], [0, h]])

    # not sure what to call this, but it's the proportion that the width of the door is narrowed to
    # when it has fully "swung in". 
    focal_point_factor = 0.1

    start_top_corner_position = (w, 0) 
    final_top_corner_position = (w * focal_point_factor, h * focal_point_factor)

    for f in range(0, frames):
        ratio = f * np.pi / (frames - 1) / 2
        tr = (
            np.cos(ratio) * (start_top_corner_position[0] - final_top_corner_position[0]) + final_top_corner_position[0],
            np.sin(ratio) * (final_top_corner_position[1] - start_top_corner_position[1]) + start_top_corner_position[1]
        )

        # Destination points (narrowing top to make a trapezoid)
        pts_dst = np.float32((
            (0, 0),  # Top-left stays
            tr,  # Top-right moved inward
            (tr[0], h - tr[1]),  # Bottom-right moved inward
            (0, h),  # Bottom-left stays
        ))

        # Get transformation matrix and warp
        matrix = cv2.getPerspectiveTransform(pts_src, pts_dst)
        # note: to adjust the image size, you can change the dsize parameter (for "door out").
        result = cv2.warpPerspective(img, matrix, (w, h))

        with tempfile.NamedTemporaryFile(suffix='.png', delete_on_close=False) as temp_image:
            cv2.imwrite(temp_image.name, result)
            pic = slide.shapes.add_picture(temp_image.name, left, top, height=Cm(10))

        pic.name = f'door-{f}'

    with open(f'{template_dir}/timing-template.xml', 'rb') as f:
        timing = etree.parse(f)

    animations = find_target(timing, "animations")

    with open (f'{template_dir}/animation-template.xml', 'rb') as f:
        animation_f = io.BytesIO(f.read())

    for i, image in enumerate(slide.element.findall(f".//{p}nvPicPr/{p}cNvPr")):
        i1 = i + 1
        # all except the first image are ..
        if i != 0:
            # .. hidden, initially
            animations.append(create_animation(animation_f, image.get("id"), 0, "exit"))
            # shown in their turn
            animations.append(create_animation(animation_f, image.get("id"), initial_delay + (i * frame_time), "entr"))
        if i1 != frames:
            # all except the last image are hidden in their turn
            animations.append(create_animation(animation_f, image.get("id"), initial_delay + (i1 * frame_time), "exit"))

        slide.element.append(timing.getroot())

    prs.save("animated-door.pptx")

    return 0


def create_animation(animation_f, image_id, delay, preset_class):
    animation = etree.parse(animation_f)

    find_target(animation, "thisDelay").attrib["delay"] = str(delay)
    find_target(animation, "spid").attrib["spid"] = image_id
    find_target(animation, "presetClass").attrib["presetClass"] = preset_class
    # TODO ok for now
    find_target(animation, "val").attrib["val"] = "hidden" if preset_class == "exit" else "visible"

    return animation.getroot()


if __name__ == "__main__":
    args = docopt(__doc__, version="0.1")
    raise SystemExit(main(
        args["<door-image-path>"],
        int(args["--frames"]),
        int(args["--frame-time"]),
        initial_delay=int(args["--initial-delay"]),
    ))
