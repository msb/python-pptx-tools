"""
While working on a Power Point presentation,
I leave image files in the project directory that have been imported into the presentation.
This script will delete any images in the project (current) directory
that are already used in the presentation
as they can be easily extracted with `extract-images.py`.

TODO: Images with halo/blurring applied seem to be different from the original image.
I think the original image is still in the presentation but I will need to bypass the model to get them.

Usage:
  tidy-images.py <path_to_file> [(-d | --dryrun)]
  tidy-images.py (-h | --help)
  tidy-images.py --version

Options:
  <path_to_file>  path to the Power Point file.
  -d --dryrun    Perform a dry run without deleting images.
  -h --help      this screen.
  --version      Show version.
"""

import hashlib
import os

from docopt import docopt
from pptx import Presentation

def tidy_images(path_to_file, dryrun=False):

  # Load the presentation and collect the hashes of all images used in the slides

  prs = Presentation(path_to_file)

  image_hashes = set()

  def get_hashes(shapes):
    """
    Recursively collect the hashes of all presentation images.
    """
    for shape in shapes:
      if hasattr(shape, "shapes"):  # `shape` is a grouping of sub-shapes
        get_hashes(shape.shapes)
      elif hasattr(shape, "image"):
        image_hashes.add(shape.image.sha1)
              
  get_hashes(prs.slides)

  # Walk the current directory and delete any images that are already presentation.
  # (the presentation is the repo for these shapes).
  for root, _, files in os.walk('.'):
    for filename in files:
      filepath = os.path.join(root, filename)
      with open(filepath, "rb") as f:
        digest = hashlib.file_digest(f, "sha1")
      if digest.hexdigest() in image_hashes:
        if dryrun:
          # Not deleted if it's a dry run.
          print(f"Would delete {filepath}")
        else:
          try:
            os.remove(filepath)
            print(f"Deleted {filepath}")
          except PermissionError as e:
            print(f"Error deleting {filepath}: {e}")
      else:
        print(f"Not in presentation. Keeping {filepath}")


if __name__ == "__main__":
  args = docopt(__doc__, version="0.1")
  tidy_images(args["<path_to_file>"], dryrun=args.get("--dryrun"))
